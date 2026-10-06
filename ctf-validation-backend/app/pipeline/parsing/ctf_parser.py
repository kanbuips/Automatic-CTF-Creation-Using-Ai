"""Parsing of the Stellantis CTF-PIS workbook (.xlsm / .xlsx).

Reads the "Front page" (identification + template version) and the "CTF & PIS" sheet
(one row per characteristic; header row located by its labels, not a fixed row number).
"""
import re
import warnings
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from app.core.exceptions import UnprocessableError

CTF_SHEET = "CTF & PIS"
FRONT_SHEET = "Front page"
HEADER_SCAN_ROWS = range(15, 60)
MAX_COL = 60
MAX_ROW = 2000

# header label (normalised) -> field; first occurrence of a label wins (the sheet repeats some)
COLUMNS = {
    "interface": "interface",
    "first level function": "function1",
    "second level function": "function2",
    "type": "type",
    "identification code": "code",
    "wording": "wording",
    "nature": "nature",
    "characteristic type": "char_type",
    "nominal value": "nominal",
    "lower specification limit": "lsl",
    "upper specification limit": "usl",
    "unit": "unit",
}
PREFIX_COLUMNS = {"it-": "it_minus", "it+": "it_plus", "priority for the compliance": "priority"}
INDICE_KEYS = {"drawing_ref": "drawing_indice", "part_ref": "part_indice"}
REQUIRED = {"code", "wording", "nominal", "lsl", "usl"}

FRONT_LABELS = {
    "part name": "part_name",
    "project code": "project",
    "reference for this part": "part_ref",
    "drawing reference": "drawing_ref",
    "version": "template_version",
    "model n°": "model_no",
    "published date": "published",
}


@dataclass
class CtfRow:
    row: int  # 1-based sheet row
    code: str
    wording: str | None = None
    interface: str | None = None
    function1: str | None = None
    function2: str | None = None
    type: str | None = None
    nature: str | None = None
    char_type: str | None = None
    nominal: float | None = None
    lsl: float | None = None
    usl: float | None = None
    it_minus: float | None = None
    it_plus: float | None = None
    unit: str | None = None
    priority: str | None = None


@dataclass
class ParsedCtf:
    front: dict[str, str | None] = field(default_factory=dict)
    rows: list[CtfRow] = field(default_factory=list)
    columns: dict[str, int] = field(default_factory=dict)  # field -> 1-based column index


def _norm(v) -> str:
    return re.sub(r"\s+", " ", str(v)).strip().lower() if v is not None else ""


def _num(v) -> float | None:
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        try:
            return float(v.strip().replace(",", "."))
        except ValueError:
            return None
    return None


def _text(v) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def _read_front(ws) -> dict[str, str | None]:
    front: dict[str, str | None] = {k: None for k in (*FRONT_LABELS.values(), "drawing_indice", "part_indice")}
    for row in ws.iter_rows(min_row=1, max_row=200, max_col=30, values_only=True):
        for i, v in enumerate(row):
            key = FRONT_LABELS.get(_norm(v).rstrip(":").strip()) if isinstance(v, str) else None
            if key and front[key] is None:
                for nxt in row[i + 1 :]:
                    if nxt not in (None, "") and not (isinstance(nxt, str) and nxt.strip().endswith(":")):
                        front[key] = str(nxt).strip()
                        break
                if key in INDICE_KEYS:  # "Indice:" sits further right on the same row
                    for j, lab in enumerate(row[i + 1 :], start=i + 1):
                        if isinstance(lab, str) and _norm(lab).rstrip(":") == "indice":
                            val = next((x for x in row[j + 1 :] if x not in (None, "")), None)
                            if val is not None and not str(val).strip().endswith(":"):
                                front[INDICE_KEYS[key]] = str(val).strip()
                            break
    return front


def _map_header(cells: tuple) -> dict[str, int]:
    cols: dict[str, int] = {}
    for idx, v in enumerate(cells, start=1):
        label = _norm(v)
        if not label:
            continue
        name = COLUMNS.get(label)
        if name is None:
            name = next((f for p, f in PREFIX_COLUMNS.items() if label.startswith(p)), None)
        if name and name not in cols:
            cols[name] = idx
    return cols


def parse_ctf(path: str | Path) -> ParsedCtf:
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001 - corrupt / wrong file type
            raise UnprocessableError(f"Cannot read workbook: {exc}") from exc
        try:
            missing = [s for s in (FRONT_SHEET, CTF_SHEET) if s not in wb.sheetnames]
            if missing:
                raise UnprocessableError(f"Not a CTF-PIS workbook: missing sheet(s) {', '.join(missing)}")
            front = _read_front(wb[FRONT_SHEET])
            grid = list(wb[CTF_SHEET].iter_rows(min_row=1, max_row=MAX_ROW, max_col=MAX_COL, values_only=True))
        finally:
            wb.close()

    header_row = next(
        (r for r in HEADER_SCAN_ROWS if r <= len(grid) and "identification code" in {_norm(v) for v in grid[r - 1]}),
        None,
    )
    if header_row is None:
        raise UnprocessableError(f"'{CTF_SHEET}' sheet: header row ('Identification code') not found")
    cols = _map_header(grid[header_row - 1])
    if REQUIRED - cols.keys():
        raise UnprocessableError(f"'{CTF_SHEET}' sheet: missing column(s) {', '.join(sorted(REQUIRED - cols.keys()))}")

    rows: list[CtfRow] = []
    for r in range(header_row + 1, len(grid) + 1):
        cells = grid[r - 1]
        code = _text(cells[cols["code"] - 1])
        if not code:
            continue
        get = lambda name: cells[cols[name] - 1] if name in cols else None  # noqa: E731
        rows.append(
            CtfRow(
                row=r,
                code=code,
                wording=_text(get("wording")),
                interface=_text(get("interface")),
                function1=_text(get("function1")),
                function2=_text(get("function2")),
                type=_text(get("type")),
                nature=_text(get("nature")),
                char_type=_text(get("char_type")),
                nominal=_num(get("nominal")),
                lsl=_num(get("lsl")),
                usl=_num(get("usl")),
                it_minus=_num(get("it_minus")),
                it_plus=_num(get("it_plus")),
                unit=_text(get("unit")),
                priority=_text(get("priority")),
            )
        )
    if not rows:
        raise UnprocessableError(f"'{CTF_SHEET}' sheet contains no characteristics")
    return ParsedCtf(front=front, rows=rows, columns=cols)
