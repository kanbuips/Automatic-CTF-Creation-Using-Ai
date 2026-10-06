"""Excel / PDF export of a report dict."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

SUMMARY_FIELDS = [
    ("Job", "job_id"),
    ("File", "filename"),
    ("Generated", "generated_at"),
    ("Part name", "part_name"),
    ("Model name", "model_name"),
    ("Template", "template_version"),
    ("Characteristics", "characteristic_count"),
    ("Verdict", "verdict"),
    ("Approval", "approval_status"),
]
FINDING_COLUMNS = ["severity", "rule", "source", "characteristic_id", "confidence", "message"]


def to_xlsx(report: dict, path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    for label, key in SUMMARY_FIELDS:
        ws.append([label, report.get(key)])
    for sev, n in report["counts"].items():
        ws.append([f"Count: {sev}", n])
    for row in ws.iter_rows(max_col=1):
        row[0].font = Font(bold=True)
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 50

    fs = wb.create_sheet("Findings")
    fs.append([c.replace("_", " ").title() for c in FINDING_COLUMNS])
    for cell in fs[1]:
        cell.font = Font(bold=True)
    for f in report["findings"]:
        fs.append([f.get(c) for c in FINDING_COLUMNS])
    fs.column_dimensions["F"].width = 80
    wb.save(path)
    return path


def to_pdf(report: dict, path: Path) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(path), pagesize=landscape(A4), title="CTF QC Report")
    story = [Paragraph("CTF QC Validation Report", styles["Title"])]
    for label, key in SUMMARY_FIELDS:
        story.append(Paragraph(f"<b>{label}:</b> {report.get(key)}", styles["Normal"]))
    story.append(Spacer(1, 12))
    counts = ", ".join(f"{k}: {v}" for k, v in report["counts"].items())
    story.append(Paragraph(f"<b>Counts</b> - {counts}", styles["Normal"]))
    story.append(Spacer(1, 12))

    body = styles["BodyText"]
    rows = [["Severity", "Rule", "Source", "Item", "Conf.", "Message"]]
    for f in report["findings"]:
        conf = "" if f.get("confidence") is None else f"{f['confidence']:.0%}"
        rows.append(
            [f["severity"], f["rule"], f["source"], f.get("characteristic_id") or "", conf, Paragraph(f["message"], body)]
        )
    table = Table(rows, repeatRows=1, colWidths=[60, 55, 55, 60, 40, 480])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
            ]
        )
    )
    story.append(table)
    doc.build(story)
    return path
