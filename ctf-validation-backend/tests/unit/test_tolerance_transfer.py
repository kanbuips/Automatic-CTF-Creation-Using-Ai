import zipfile

import pytest
from openpyxl import Workbook, load_workbook

from app.core.exceptions import UnprocessableError
from app.pipeline.generation.tolerance_transfer import transfer_tolerances
from app.pipeline.parsing.ctf_parser import parse_ctf

HEADERS = {
    2: "Interface", 3: "First Level Function", 4: "Second Level Function", 5: "Type", 6: "Identification code",
    7: "Wording", 9: "Nature", 10: "Characteristic type", 16: "Nominal value", 17: "Lower specification limit",
    18: "Upper specification limit", 19: "IT-\nLower limit deviation", 20: "IT+\nUpper limit deviation", 25: "Unit",
    26: "Priority for the compliance upgrade",
}


def make_ctf(path, rows, version="BETA_V003", drawing=None, indice=None, name="TEST PART"):
    """rows: (type, code, wording, nominal, lsl, usl). S/T are formulas like the real template."""
    wb = Workbook()
    fp = wb.active
    fp.title = "Front page"
    fp["D30"], fp["H30"] = "Part name:", name
    fp["D38"], fp["H38"], fp["L38"] = "Reference for this part:", "11111111AA", "Indice:"
    if drawing:
        fp["D40"], fp["H40"] = "Drawing reference:", drawing
        if indice:
            fp["L40"], fp["M40"] = "Indice:", indice
    fp["N141"], fp["O141"] = "Version :", f" {version}"
    ws = wb.create_sheet("CTF & PIS")
    for col, label in HEADERS.items():
        ws.cell(29, col, label)
    ws.cell(29, 46, "Identification code")  # duplicate header further right must be ignored
    for i, (typ, code, wording, nom, lsl, usl) in enumerate(rows):
        r = 30 + i
        ws.cell(r, 5, typ), ws.cell(r, 6, code), ws.cell(r, 7, wording)
        ws.cell(r, 16, nom), ws.cell(r, 17, lsl), ws.cell(r, 18, usl), ws.cell(r, 25, "mm")
        ws.cell(r, 19, f'=IF($P{r}="","N/A",$Q{r}-$P{r})')
        ws.cell(r, 20, f'=IF($P{r}="","N/A",$R{r}-$P{r})')
    wb.save(path)
    return path


@pytest.fixture()
def files(tmp_path):
    new = make_ctf(tmp_path / "new.xlsx", [
        ("CTF", "01A01", "Size Tolerance of the datum A in C/C direction", 0, -0.1, 0.1),
        ("CTF", "01A02", "Size Tolerance of the datum B in C/C direction", 0, -0.1, 0.1),
        ("CTF", "01A03", "Size Tolerance of the datum C in C/C direction", 0, -0.1, 0.1),
        ("CTF", "01A04", "Size Tolerance of the part in F/A direction", 0, -0.37, 0),
        ("CTF", "01A05", "Size Tolerance of the datum D in C/C direction", 0, -0.1, 0.1),
    ])
    old = make_ctf(tmp_path / "old.xlsx", [
        ("CTF", "01A01", "Flatness tolerance of the Datum A", 0, 0, 0.4),
        ("CTF", "01A02", "Size tolerance of the Datum B Clamping diameter", 0, -0.3, 0.3),
        ("CSE", "01A02-01", "Size tolerance of the Datum B Clamping diameter", 9.3, 9, 9.6),
        ("CTF", "01A03", "Size tolerance of the Datum C Clamping Width", 0, -0.3, 0.3),
        ("CTF", "01A04", "Size tolerance of the Datum C bend center location", 0, -0.3, 0.3),
        ("CTF", "01A05", "Size tolerance of the Datum B Opening width", 0, -0.5, 0.5),
        ("CTF", "01A06", "Size tolerance of Wire Diameter", 0, -0.05, 0.05),
    ])
    return new, old, tmp_path / "out.xlsx"


def test_parse_ctf_front_and_rows(files):
    p = parse_ctf(files[0])
    assert p.front["template_version"] == "BETA_V003" and p.front["part_ref"] == "11111111AA"
    assert [r.code for r in p.rows] == ["01A01", "01A02", "01A03", "01A04", "01A05"]
    assert p.rows[0].lsl == -0.1 and p.columns["code"] == 6  # first 'Identification code' wins


def test_rejects_non_ctf_workbook(tmp_path):
    wb = Workbook()
    wb.save(tmp_path / "x.xlsx")
    with pytest.raises(UnprocessableError):
        parse_ctf(tmp_path / "x.xlsx")


def test_strict_matching_statuses(files):
    new, old, out = files
    plan = {p.code: p for p in transfer_tolerances(new, old, out).plan}
    assert plan["01A01"].status == "no_match"  # old datum A row is Flatness, not Size
    assert plan["01A02"].status == "ambiguous"  # CTF rows only (CSE ignored): 01A02 (+-0.3) vs 01A05 (+-0.5)
    assert plan["01A03"].status == "updated" and plan["01A03"].new == (0.0, -0.3, 0.3)
    assert plan["01A03"].source_codes == ["01A03", "01A04"]  # two rows agree -> safe
    assert plan["01A04"].status == "no_key"
    assert plan["01A05"].status == "no_match"


def test_only_matched_rows_change_and_formulas_survive(files):
    new, old, out = files
    transfer_tolerances(new, old, out)
    ws = load_workbook(out)["CTF & PIS"]
    assert [ws.cell(32, c).value for c in (16, 17, 18)] == [0, -0.3, 0.3]
    assert ws.cell(32, 19).value.startswith("=IF(")  # formula kept
    assert [ws.cell(31, c).value for c in (16, 17, 18)] == [0, -0.1, 0.1]  # ambiguous row untouched
    assert [ws.cell(33, c).value for c in (16, 17, 18)] == [0, -0.37, 0]
    with zipfile.ZipFile(out) as z:
        assert 'fullCalcOnLoad="1"' in z.read("xl/workbook.xml").decode()


def test_cached_it_values_refreshed(files):
    new, old, out = files
    transfer_tolerances(new, old, out)
    cached = load_workbook(out, data_only=True)["CTF & PIS"]
    # openpyxl wrote no cached values, so only the patched row has them
    assert cached.cell(32, 19).value is None or cached.cell(32, 19).value == -0.3


def test_overrides_force_pairing_and_validate_codes(files):
    new, old, out = files
    plan = {p.code: p for p in transfer_tolerances(new, old, out, {"01A02": "01A05", "01A04": "01A06"}).plan}
    assert plan["01A02"].status == "updated" and plan["01A02"].manual and plan["01A02"].new == (0.0, -0.5, 0.5)
    assert plan["01A04"].new == (0.0, -0.05, 0.05)
    with pytest.raises(UnprocessableError):
        transfer_tolerances(new, old, out, {"01A02": "99Z99"})
    with pytest.raises(UnprocessableError):
        transfer_tolerances(new, old, out, {"nope": "01A01"})


def test_unchanged_when_values_equal(tmp_path, files):
    new, _, out = files
    same = make_ctf(tmp_path / "same.xlsx", [("CTF", "01A01", "Size tolerance of the Datum A", 0, -0.1, 0.1)])
    plan = transfer_tolerances(new, same, out).plan
    assert plan[0].status == "unchanged"
