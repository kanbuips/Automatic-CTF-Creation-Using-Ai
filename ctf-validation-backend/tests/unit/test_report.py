from openpyxl import load_workbook

from app.core.config import get_settings
from app.pipeline.parsing.jtxml_parser import parse_jtxml_bytes
from app.pipeline.reporting.exporters import to_pdf, to_xlsx
from app.pipeline.reporting.report_builder import build_report, compute_verdict
from app.pipeline.rules.base import RuleFinding
from app.pipeline.rules.engine import run_rules


def test_verdicts():
    assert compute_verdict({}) == "pass"
    assert compute_verdict({"warning": 1}) == "review"
    assert compute_verdict({"warning": 1, "error": 1}) == "fail"
    assert compute_verdict({"critical": 1}) == "fail"


def test_build_report_merges_sources(bad_bytes):
    parsed = parse_jtxml_bytes(bad_bytes)
    findings = run_rules(parsed, get_settings()) + [RuleFinding("ANM001", "warning", "odd", "C3", source="anomaly", confidence=0.8)]
    r = build_report({"job_id": "j", "filename": "f.jtxml"}, parsed, findings)
    assert r["verdict"] == "fail" and r["counts"]["total"] == len(findings)
    assert r["by_source"]["anomaly"] == 1 and r["anomalies"] == ["C3"]
    assert r["findings"][0]["severity"] == "critical"


def test_exports(tmp_path, good_bytes, bad_bytes):
    parsed = parse_jtxml_bytes(bad_bytes)
    report = build_report({"job_id": "j", "filename": "f"}, parsed, run_rules(parsed, get_settings()))
    xlsx = to_xlsx(report, tmp_path / "r.xlsx")
    wb = load_workbook(xlsx)
    assert wb.sheetnames == ["Summary", "Findings"] and wb["Findings"].max_row == len(report["findings"]) + 1
    pdf = to_pdf(report, tmp_path / "r.pdf")
    assert pdf.read_bytes().startswith(b"%PDF")
